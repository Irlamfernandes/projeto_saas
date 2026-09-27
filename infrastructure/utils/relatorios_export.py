"""Exportação dos relatórios em PDF e Excel a partir das mesmas definições de seção."""
from dataclasses import dataclass
from functools import partial
from decimal import Decimal
from typing import Callable
import openpyxl
from django.utils import timezone
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from infrastructure.models.produto import sigla
from infrastructure.utils.formatacao import numero, reais
from infrastructure.utils.pdf_marca import cabecalho_da_loja

FORMATO_MOEDA_EXCEL = '#,##0.00'
COR_RODAPE = colors.HexColor('#9ca3af')
PREFIXOS_DE_FORMULA = ('=', '+', '-', '@')


@dataclass(frozen=True)
class Secao:
    titulo: str
    aba: str
    cabecalho: tuple
    linhas: Callable
    larguras_pdf: tuple = (300, 212)


RESUMO = Secao(
    titulo="Resumo Executivo e Indicadores",
    aba="Resumo Executivo",
    cabecalho=("Indicador", "Valor"),
    linhas=lambda r: [
        ("Faturamento Total", r.total_geral_vendas),
        ("Ticket Médio", r.ticket_medio),
        ("Total a Receber (Fiado)", r.total_fiado_pendente),
        ("Crédito de Clientes", r.total_credito_clientes),
    ],
)
TEMPORAL = Secao(
    titulo="Desempenho por Ciclo Temporal",
    aba="Desempenho Temporal",
    cabecalho=("Período", "Faturamento"),
    linhas=lambda r: [
        ("Hoje", r.vendas_hoje_valor),
        ("Últimos 7 dias", r.vendas_semana_valor),
        ("Este Mês", r.vendas_mes_valor),
        ("Este Ano", r.vendas_ano_valor),
    ],
)
DEVENDO = Secao(
    titulo="Relatório de Clientes Devendo (Fiado)",
    aba="Clientes Devendo",
    cabecalho=("Cliente", "Valor Devido"),
    linhas=lambda r: [(c.nome, c.saldo_calculado) for c in r.clientes_devendo],
)
CREDITO = Secao(
    titulo="Relatório de Clientes com Crédito",
    aba="Clientes com Crédito",
    cabecalho=("Cliente", "Crédito Disponível"),
    linhas=lambda r: [(c.nome, c.valor_credito) for c in r.clientes_credito],
)
ESTOQUE = Secao(
    titulo="Inventário de Estoque Completo",
    aba="Estoque",
    cabecalho=("Produto", "Preço de Venda", "Qtd em Estoque", "Unidade"),
    linhas=lambda r: [
        (p.nome, p.preco_venda, _quantidade_numerica(p.quantidade_estoque), sigla(p.unidade)) for p in r.estoque_ordenado
    ],
    larguras_pdf=(220, 110, 110, 72),
)

SECOES_POR_TIPO = {
    'resumo': (RESUMO,),
    'temporal': (TEMPORAL,),
    'devendo': (DEVENDO,),
    'credito': (CREDITO,),
    'estoque': (ESTOQUE,),
    'geral': (RESUMO, TEMPORAL, DEVENDO, CREDITO, ESTOQUE),
}
TIPOS_DE_RELATORIO = frozenset(SECOES_POR_TIPO)

_ESTILO_TABELA = TableStyle([
    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f3f4f6')),
    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
])


def gerar_relatorio_pdf(tipo, relatorio, destino, loja=None, storage=None, assinatura=''):
    """`loja`: põe o cabeçalho com nome e marca (o superusuário, sem loja, exporta sem cabeçalho).
    `assinatura`: texto pequeno no canto direito do rodapé de cada página (quem fez o sistema)."""
    titulo = ParagraphStyle(
        'TituloRelatorio', parent=getSampleStyleSheet()['Heading1'],
        fontSize=16, textColor=colors.HexColor('#1f2937'), spaceAfter=10,
    )
    story = cabecalho_da_loja(loja, "Relatório", storage) if loja else []
    for secao in SECOES_POR_TIPO[tipo]:
        story += _secao_pdf(secao, relatorio, titulo)

    rodape = partial(_rodape_da_pagina, gerado_em=timezone.localtime(), assinatura=assinatura)
    SimpleDocTemplate(destino, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=40).build(
        story, onFirstPage=rodape, onLaterPages=rodape,
    )


def _rodape_da_pagina(canvas, documento, *, gerado_em, assinatura):
    """Linha discreta no pé da página: quando foi gerado, qual página e quem fez o sistema."""
    canvas.saveState()
    canvas.setFont('Helvetica', 7.5)
    canvas.setFillColor(COR_RODAPE)
    canvas.drawString(documento.leftMargin, 20, f"Gerado em {gerado_em:%d/%m/%Y} às {gerado_em:%H:%M} · Página {documento.page}")
    if assinatura:
        canvas.drawRightString(documento.pagesize[0] - documento.rightMargin, 20, assinatura)
    canvas.restoreState()


def gerar_relatorio_excel(tipo, relatorio, destino):
    planilha = openpyxl.Workbook()
    planilha.remove(planilha.active)
    for secao in SECOES_POR_TIPO[tipo]:
        _preencher_aba(planilha.create_sheet(secao.aba), secao, relatorio)
    planilha.save(destino)


def _secao_pdf(secao, relatorio, estilo_titulo):
    dados = [list(secao.cabecalho)] + [[_texto_pdf(valor) for valor in linha] for linha in secao.linhas(relatorio)]
    tabela = Table(dados, colWidths=list(secao.larguras_pdf))
    tabela.setStyle(_ESTILO_TABELA)
    return [Paragraph(f"<b>{secao.titulo}</b>", estilo_titulo), Spacer(1, 10), tabela, Spacer(1, 15)]


def _quantidade_numerica(valor):
    """Quantidade como número comum (3 ou 2.5): Decimal nas planilhas e PDFs é tratado como dinheiro."""
    return int(valor) if valor == valor.to_integral_value() else float(valor)


def _texto_pdf(valor):
    if isinstance(valor, Decimal):
        return reais(valor)
    if isinstance(valor, float):
        return numero(valor)
    return str(valor)


def _preencher_aba(aba, secao, relatorio):
    aba.append(secao.cabecalho)
    for linha in secao.linhas(relatorio):
        aba.append(linha)
        for celula in aba[aba.max_row]:
            _formatar_celula(celula)


def _formatar_celula(celula):
    if isinstance(celula.value, Decimal):
        celula.number_format = FORMATO_MOEDA_EXCEL
    elif isinstance(celula.value, str) and celula.value.startswith(PREFIXOS_DE_FORMULA):
        # Nome digitado pelo usuário nunca pode virar fórmula ao abrir a planilha.
        celula.data_type = 's'
